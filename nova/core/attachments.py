from __future__ import annotations

import base64
import io
import json
import uuid
import zipfile
from pathlib import Path

MAX_BYTES = 8 * 1024 * 1024
# Text extraction cap per attachment. The full text is only pasted inline when it
# is small enough (see rag.index_above_chars); larger documents are chunk-indexed
# and answered from retrieved fragments, so a generous cap here is safe.
MAX_TEXT = 200000
TEXT_TYPES = {'.txt', '.md', '.py', '.js', '.ts', '.tsx', '.jsx', '.json', '.yaml', '.yml',
              '.csv', '.html', '.css', '.xml', '.sql', '.log', '.toml', '.ini', '.java', '.c', '.cpp', '.h', '.rs', '.go'}


class AttachmentStore:
    """Local library with opaque identifiers and bounded, passive extraction."""

    def __init__(self, root: Path):
        self.root = root.resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def _path(self, key: str, suffix: str) -> Path:
        if len(key) != 32 or any(c not in '0123456789abcdef' for c in key):
            raise ValueError('Archivo no encontrado.')
        path = self.root / (key + suffix)
        if path.is_symlink() or path.resolve().parent != self.root:
            raise ValueError('Archivo no disponible.')
        return path

    def add(self, name: str, encoded: str) -> dict:
        if len(encoded) > (MAX_BYTES * 4 // 3 + 4):
            raise ValueError('El archivo supera el límite de 8 MB.')
        try:
            raw = base64.b64decode(encoded, validate=True)
        except ValueError:
            raise ValueError('No hemos podido leer el archivo.') from None
        if not raw or len(raw) > MAX_BYTES:
            raise ValueError('El archivo está vacío o supera los 8 MB.')
        name = name.replace('\\', '/').rsplit('/', 1)[-1][:160]
        suffix = Path(name).suffix.lower()
        text, image = '', ''
        if suffix in {'.png', '.jpg', '.jpeg', '.webp', '.gif', '.bmp'}:
            from PIL import Image, ImageOps

            with Image.open(io.BytesIO(raw)) as picture:
                if picture.width * picture.height > 25_000_000:
                    raise ValueError('La imagen es demasiado grande. Usa una copia de menor resolución.')
                picture = ImageOps.exif_transpose(picture).convert('RGB')
                picture.thumbnail((1600, 1600))
                output = io.BytesIO()
                picture.save(output, 'JPEG', quality=85)
            raw = output.getvalue()
            image = 'data:image/jpeg;base64,' + base64.b64encode(raw).decode('ascii')
        elif suffix == '.pdf':
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError('Este PDF está protegido. Adjunta una copia sin contraseña.')
            parts = []
            for page in reader.pages[:40]:
                content = page.get_contents()
                if content is not None and len(content.get_data()) > 4 * 1024 * 1024:
                    raise ValueError('Este PDF es demasiado complejo. Adjunta menos páginas.')
                parts.append((page.extract_text() or '')[:MAX_TEXT])
                if sum(map(len, parts)) >= MAX_TEXT:
                    break
            text = '\n'.join(parts)[:MAX_TEXT]
            if not text.strip():
                raise ValueError('Este PDF no contiene texto legible. Puedes adjuntar una imagen de la página.')
        elif suffix == '.docx':
            from docx import Document

            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 32 * 1024 * 1024:
                    raise ValueError('El documento descomprimido es demasiado grande.')
            doc = Document(io.BytesIO(raw))
            text = '\n'.join([p.text for p in doc.paragraphs] +
                             [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])[:MAX_TEXT]
        elif suffix in TEXT_TYPES:
            text = raw.decode('utf-8-sig')[:MAX_TEXT]
        else:
            raise ValueError('Puedes adjuntar imágenes, PDF, DOCX, texto y archivos de código.')
        key = uuid.uuid4().hex
        metadata = dict(id=key, name=name, size=len(raw), kind='image' if image else 'document',
                        text=text, image=image, truncated=len(text) >= MAX_TEXT)
        self._path(key, '.bin').write_bytes(raw)
        self._path(key, '.json').write_text(json.dumps(metadata, ensure_ascii=False), encoding='utf-8')
        return self.public(metadata)

    def get(self, key: str) -> dict:
        try:
            return json.loads(self._path(key, '.json').read_text(encoding='utf-8'))
        except (OSError, ValueError):
            raise ValueError('Archivo no encontrado en la biblioteca.') from None

    @staticmethod
    def public(item: dict) -> dict:
        return {k: item[k] for k in ('id', 'name', 'size', 'kind', 'truncated')}

    def list(self) -> list[dict]:
        return [self.public(self.get(p.stem)) for p in sorted(self.root.glob('*.json'), key=lambda p: p.stat().st_mtime, reverse=True)]

    def delete(self, key: str):
        self.get(key)
        self._path(key, '.json').unlink()
        self._path(key, '.bin').unlink(missing_ok=True)
